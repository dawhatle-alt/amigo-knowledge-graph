#!/usr/bin/env python3
"""
One-shot converter: AMIGO_Checklist_V22_*.xlsx -> Obsidian vault notes.

Run once to seed the vault. After that the VAULT is the source of truth;
re-running will overwrite generated notes (checklist/, runbook/, components/,
procedures/, _meta/) but never touches hand-maintained notes (rules/, versions/,
paths/, CLAUDE.md, README.md).

Usage: python3 tools/convert_xlsx.py <xlsx> <vault_root>
"""
import re, sys, os, json
from collections import defaultdict
import openpyxl

XLSX, ROOT = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(XLSX, data_only=True)
SRC = os.path.basename(XLSX)

def slug(s, n=48):
    s = re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")
    return s[:n].rstrip("-")

def clean(s):
    if s is None: return ""
    return re.sub(r"[ \t\xa0]+\n", "\n", str(s).replace("\xa0", " ")).strip()

def yaml_list(xs):
    return "[" + ", ".join(json.dumps(x) for x in xs) + "]"

def write(path, text):
    full = os.path.join(ROOT, path)
    os.makedirs(os.path.dirname(full), exist_ok=True)
    with open(full, "w", encoding="utf-8") as f:
        f.write(text)

source_map = []

# ---------- data validations (dropdowns) ----------
def dropdowns(ws):
    """cell -> list of option sets"""
    out = defaultdict(list)
    for dv in ws.data_validations.dataValidation:
        if not dv.formula1: continue
        opts = [o.strip() for o in dv.formula1.strip('"').split(",") if o.strip() and o.strip().lower() != "select one"]
        if not opts: continue
        for rng in str(dv.sqref).split():
            if ":" in rng:
                a, b = rng.split(":")
                col = re.match(r"[A-Z]+", a).group()
                r1, r2 = int(a[len(col):]), int(b[len(col):])
                for r in range(r1, r2 + 1): out[f"{col}{r}"].append(opts)
            else:
                out[rng].append(opts)
    return out

def answer_type(opts):
    if not opts: return "free-text"
    s = set(o.lower() for o in opts)
    if s <= {"yes", "no", "na", "na - doing in-place", "not applicable"}: return "yes-no"
    if s <= {"applicable", "not applicable"}: return "applicable"
    if s <= {"done", "not applicable"}: return "done"
    if s == {"acknowledged"}: return "acknowledge"
    if s <= {"same", "different"}: return "same-different"
    if s <= {"in-place", "migration"}: return "in-place-migration"
    if any(re.match(r"9\.0\.\d+", o) for o in opts): return "version"
    return "select"

# rows whose prose is overridden by a hand-maintained rule note
OVERRIDES = {
    ("em", 43): {"status": "superseded", "superseded_by": "direct-upgrade-requires-9020",
                 "note": "Source text says 'at least 9.0.19'. This contradicts the FROM-version dropdown on row 11 "
                         "(9.0.20, 9.0.21 only) and the AMIGO domain rule. The rule note is authoritative."},
}

# keyword -> linked rule / component
RULE_LINKS = [
    (r"compatibility mode", "compatibility-mode-irreversible"),
    (r"postgres", "postgres-not-upgraded-in-place"),
    (r"same or higher version", "em-before-server"),
    (r"different machine|migrat", "amigo-scope"),
    (r"severity 1|sev.?1", "amigo-scope"),
    (r"z/os", "amigo-scope"),
    (r"upgrade directly to 9\.0\.22|at least 9\.0\.19", "direct-upgrade-requires-9020"),
    (r"high availability|distributed", "ha-distributed-upgrade-order"),
    (r"authoriz", "authorizations-to-roles"),
]
COMPONENT_NAMES = {
    "Batch Impact Manager": "bim", "Forecast": "forecast", "Workflow Insights": "workflow-insights",
    "Workload Archiving": "workload-archiving", "Automation API": "aapi", "Self Service": "self-service",
    "Workload Change Manager": "wcm", "Application Intergrator": "application-integrator",
    "Application Integrator": "application-integrator", "Manage File Transfer": "mft",
}
def links_for(text):
    t = text.lower()
    rules = sorted({r for pat, r in RULE_LINKS if re.search(pat, t)})
    comps = sorted({c for name, c in COMPONENT_NAMES.items() if name.lower() in t})
    return rules, comps

# ---------- checklist sheets ----------
def convert_checklist(sheet, comp, comp_label):
    ws = wb[sheet]
    dv = dropdowns(ws)
    section = None
    notes = []
    last_q = None
    for r in range(10, ws.max_row + 1):
        a = clean(ws.cell(r, 1).value)
        if not a or a.startswith("DISCLAIMER") or a.startswith("BMC Software"): continue
        others = [clean(ws.cell(r, c).value) for c in range(2, 8)]
        if a.startswith("AMIGO") and not any(others):
            section = re.sub(r"^AMIGO\s*-\s*", "", a).replace("Conrol-M", "Control-M").replace("Contol-M", "Control-M")
            continue
        if a.startswith("*") and last_q:
            item = a.lstrip("* ").strip()
            last_q["subitems"].append({"row": r, "text": item, "customer": others[1], "tsa": others[5],
                                       "tsa_options": [o for s in dv.get(f"E{r}", []) for o in s]})
            continue
        q = {"row": r, "section": section, "question": a, "customer_guidance": others[1],
             "tsa_guidance": others[5], "cust_opts": dv.get(f"B{r}", []), "tsa_opts": dv.get(f"E{r}", []),
             "subitems": []}
        notes.append(q); last_q = q

    idx_lines = []
    for q in notes:
        short = re.split(r"[\n?.]", q["question"])[0]
        name = f"{comp}-{q['row']:03d}-{slug(short, 40)}"
        opts = q["cust_opts"]
        primary = opts[0] if opts else []
        atype = answer_type(primary)
        rules, comps = links_for(q["question"] + " " + q["customer_guidance"] + " " + q["tsa_guidance"])
        for s in q["subitems"]:
            _, c2 = links_for(s["text"]); comps = sorted(set(comps) | set(c2))
        ov = OVERRIDES.get((comp, q["row"]), {})
        blocking = bool(rules & {"direct-upgrade-requires-9020", "amigo-scope", "em-before-server"}) if False else \
                   any(x in rules for x in ["direct-upgrade-requires-9020", "amigo-scope", "em-before-server"])
        fm = [
            "---", f"name: {name}", "type: checklist-item", f"component: {comp}",
            f"section: {json.dumps(q['section'] or '')}", f"answer_type: {atype}",
            f"answer_options: {yaml_list(primary)}",
        ]
        if len(opts) > 1:
            fm.append(f"answer_options_alt: {yaml_list([o for s in opts[1:] for o in s])}")
            fm.append("answer_options_source_conflict: true")
        tsa_opts = [o for s in q["tsa_opts"] for o in s]
        if tsa_opts: fm.append(f"tsa_answer_options: {yaml_list(sorted(set(tsa_opts)))}")
        fm += [f"blocking: {'true' if blocking else 'false'}",
               f"related_rules: {yaml_list(rules)}", f"related_components: {yaml_list(comps)}",
               f"status: {ov.get('status', 'active')}"]
        if "superseded_by" in ov: fm.append(f"superseded_by: {ov['superseded_by']}")
        fm += [f"source_sheet: {json.dumps(sheet.strip())}", f"source_row: {q['row']}", f"source_file: {SRC}", "---", ""]
        body = [f"# {q['question'].splitlines()[0]}", ""]
        if "note" in ov:
            body += [f"> [!warning] Superseded — see [[{ov['superseded_by']}]]", f"> {ov['note']}", ""]
        body += ["## Question", "", q["question"], ""]
        if q["customer_guidance"]: body += ["## Customer guidance", "", q["customer_guidance"], ""]
        if q["tsa_guidance"]: body += ["## TSA guidance", "", q["tsa_guidance"], ""]
        if q["subitems"]:
            body += ["## Items", ""]
            for s in q["subitems"]:
                _, c2 = links_for(s["text"])
                link = f" → [[{c2[0]}]]" if c2 else ""
                body.append(f"- **{s['text']}**{link}")
                if s["customer"]: body.append(f"  - Customer: {s['customer'].replace(chr(10), ' ')}")
                if s["tsa"]: body.append(f"  - TSA: {s['tsa'].replace(chr(10), ' ')}")
            body.append("")
        if rules or comps:
            body += ["## Related", ""]
            body += [f"- Rule: [[{x}]]" for x in rules] + [f"- Component: [[{x}]]" for x in comps] + [""]
        write(f"checklist/{comp}/{name}.md", "\n".join(fm + body))
        if comp == "em" and q["row"] == 29: globals()["ADDON_NOTE"] = name
        idx_lines.append(f"- [[{name}]] — {short.strip()}" + ("  ⚠ superseded" if ov else ""))
        source_map.append((sheet.strip(), q["row"], f"checklist/{comp}/{name}.md"))

    write(f"checklist/{comp}/checklist-{comp}-index.md", "\n".join([
        "---", f"name: checklist-{comp}-index", "type: index", f"component: {comp}", "---", "",
        f"# {comp_label} AMIGO Starter Checklist — index", "",
        "Ordered as in the source spreadsheet. Section headers are preserved in each note's `section` field.", ""] + idx_lines + [""]))

convert_checklist("AMIGO EM Checklist V22", "em", "Control-M/Enterprise Manager")
convert_checklist("AMIGO Server Checklist V22 ", "server", "Control-M/Server")

# ---------- runbook sheets ----------
TOPO_PATTERNS = [
    (r"^For Standalone", "standalone", "Standalone"),
    (r"with Distributed configured with High Availability", "distributed-ha", "Distributed + High Availability"),
    (r"with Distributed", "distributed", "Distributed"),
    (r"configured with High Availability", "ha", "High Availability"),
]
def convert_runbook(sheet, comp, comp_label):
    ws = wb[sheet]
    phases = []; cur = None; topo = None; seen = set()
    for r in range(2, ws.max_row + 1):
        a = clean(ws.cell(r, 1).value)
        if not a or a.startswith("DISCLAIMER") or a.startswith("BMC Software"):
            if topo and not a: topo = None
            continue
        if clean(ws.cell(r, 5).value) == "Start":
            cur = {"row": r, "title": a, "items": [], "topologies": []}; phases.append(cur); topo = None; continue
        if cur is None: continue
        m = next(((k, lbl) for pat, k, lbl in TOPO_PATTERNS if re.search(pat, a)), None)
        if m and a.startswith("For "):
            head, _, rest = a.partition("\n")
            topo = {"key": m[0], "label": m[1], "row": r, "heading": head.rstrip(":"), "steps": [], "notes": []}
            cur["topologies"].append(topo)
            for line in rest.splitlines():
                if line.strip(): _push_step(topo, line)
            continue
        if topo and not re.match(r"^\d+\.", a):
            topo = None  # numbered block ended without a blank row
        if topo:
            key = re.sub(r"\s+", " ", a)
            if key in seen: continue
            seen.add(key)
            for line in a.splitlines():
                if line.strip(): _push_step(topo, line)
            continue
        raw = str(ws.cell(r, 1).value)
        if raw.startswith("   "):  # indented doc reference
            if cur["items"]: cur["items"][-1]["refs"].append(a.replace("For Control-M 9.0.22:", "").strip())
        else:
            cur["items"].append({"row": r, "text": a, "refs": []})
    return phases

def _push_step(topo, line):
    s = line.strip()
    if re.match(r"^\d+\.\s", s): topo["steps"].append(re.sub(r"^\d+\.\s+", "", s))
    elif s.lower().startswith("note:"): topo["notes"].append(s)
    elif topo["steps"]: topo["steps"][-1] += f" — {s}"
    else: topo["notes"].append(s)

def write_runbook(sheet, comp, comp_label):
    phases = convert_runbook(sheet, comp, comp_label)
    phase_names = []
    for i, p in enumerate(phases, 1):
        pname = f"{comp}-phase-{i}-{slug(p['title'], 36)}"
        phase_names.append((pname, p["title"]))
        rules, _ = links_for(" ".join(it["text"] for it in p["items"]))
        fm = ["---", f"name: {pname}", "type: runbook-phase", f"component: {comp}", f"phase_order: {i}",
              f"phase_title: {json.dumps(p['title'])}", "target_version: 9.0.22",
              f"related_rules: {yaml_list(rules)}", f"source_sheet: {json.dumps(sheet)}", f"source_row: {p['row']}",
              f"source_file: {SRC}", "---", "", f"# {comp_label} — Phase {i}: {p['title']}", ""]
        body = []
        for it in p["items"]:
            body.append(f"- [ ] {it['text'].replace(chr(10), '  ' + chr(10) + '      ')}")
            for ref in it["refs"]: body.append(f"    - Doc: {ref} (9.0.22)")
        if p["topologies"]:
            body += ["", "## Topology-specific sequence", "",
                     "Load exactly one of these based on the customer's topology answer:", ""]
            for t in p["topologies"]:
                tname = f"{comp}-topology-{t['key']}"
                body.append(f"- [[{tname}]] — {t['label']}")
                tfm = ["---", f"name: {tname}", "type: runbook-topology", f"component: {comp}",
                       f"topology: {t['key']}", "target_version: 9.0.22", f"parent_phase: {pname}",
                       f"related_rules: {yaml_list(['ha-distributed-upgrade-order', 'postgres-not-upgraded-in-place'] if t['key'] != 'standalone' else ['postgres-not-upgraded-in-place'])}",
                       f"source_sheet: {json.dumps(sheet)}", f"source_row: {t['row']}", f"source_file: {SRC}", "---", "",
                       f"# {comp_label} upgrade sequence — {t['label']}", "", f"_{t['heading']}_", ""]
                if t["notes"]: tfm += [f"> [!note] {n}" for n in t["notes"]] + [""]
                tfm += ["## Steps", ""] + [f"{n}. {s}" for n, s in enumerate(t["steps"], 1)] + ["",
                        "## Related", "", f"- Phase: [[{pname}]]", "- Rule: [[ha-distributed-upgrade-order]]" if t["key"] != "standalone" else "",
                        "- Rule: [[postgres-not-upgraded-in-place]]", ""]
                write(f"runbook/{comp}/{tname}.md", "\n".join(x for x in tfm if x is not None))
                source_map.append((sheet, t["row"], f"runbook/{comp}/{tname}.md"))
        if rules: body += ["", "## Related rules", ""] + [f"- [[{x}]]" for x in rules]
        write(f"runbook/{comp}/{pname}.md", "\n".join(fm + body + [""]))
        source_map.append((sheet, p["row"], f"runbook/{comp}/{pname}.md"))
    write(f"runbook/{comp}/runbook-{comp}-index.md", "\n".join(
        ["---", f"name: runbook-{comp}-index", "type: index", f"component: {comp}", "---", "",
         f"# {comp_label} sample upgrade plan (9.0.22) — phases", "", "Execute in order."] +
        [f"{i}. [[{n}]] — {t}" for i, (n, t) in enumerate(phase_names, 1)] + [""]))

write_runbook("EM V22 Upgrade", "em", "Control-M/Enterprise Manager")
write_runbook("Server V22 Upgrade", "server", "Control-M/Server")

# ---------- components (from EM checklist add-on rows) ----------
ws = wb["AMIGO EM Checklist V22"]
for r in range(30, 39):
    text = clean(ws.cell(r, 1).value).lstrip("* ")
    cust, tsa = clean(ws.cell(r, 3).value), clean(ws.cell(r, 7).value)
    _, comps = links_for(text)
    if not comps: continue
    key = comps[0]
    kas = sorted(set(re.findall(r"\b0003\d{5}|\b0004\d{5}|\b0002\d{5}", cust + " " + tsa)))
    fm = ["---", f"name: {key}", "type: component", f"display_name: {json.dumps(text)}", f"knowledge_articles: {yaml_list(kas)}",
          "source_sheet: \"AMIGO EM Checklist V22\"", f"source_row: {r}", f"source_file: {SRC}", "---", "",
          f"# {text}", "", "Add-on component to check during the AMIGO EM discussion.", ""]
    if cust: fm += ["## Customer guidance", "", cust, ""]
    if tsa: fm += ["## TSA guidance", "", tsa, ""]
    fm += ["## Related", "", f"- Checklist: [[{ADDON_NOTE}]]", ""]
    write(f"components/{key}.md", "\n".join(fm))
    source_map.append(("AMIGO EM Checklist V22", r, f"components/{key}.md"))

# ---------- procedures ----------
ws = wb["AMIGO Starter Procedures"]
lines = ["---", "name: amigo-starter-case-procedure", "type: procedure", "audience: tsa",
         "source_sheet: \"AMIGO Starter Procedures\"", "source_row: 2", f"source_file: {SRC}", "---", "",
         "# AMIGO Starter case procedure", "", "TSA-side case handling steps, in order.", "",
         "| # | Action | Actionable by |", "|---|---|---|"]
n = 0
for r in range(3, ws.max_row + 1):
    a, who = clean(ws.cell(r, 1).value), clean(ws.cell(r, 3).value)
    if not a or a.startswith("DISCLAIMER") or a.startswith("BMC"): continue
    n += 1; lines.append(f"| {n} | {a} | {who.replace('Cusomter', 'Customer')} |")
lines += ["", "## Related", "", "- [[amigo-scope]]", "- [[checklist-em-index]]", "- [[checklist-server-index]]", ""]
write("procedures/amigo-starter-case-procedure.md", "\n".join(lines))
source_map.append(("AMIGO Starter Procedures", 2, "procedures/amigo-starter-case-procedure.md"))

# ---------- source map ----------
sm = ["---", "name: source-map", "type: meta", f"source_file: {SRC}", "---", "",
      f"# Source map — `{SRC}`", "", "Generated by `tools/convert_xlsx.py`. Maps spreadsheet rows to vault notes.", "",
      "| Sheet | Row | Note |", "|---|---|---|"]
for s, r, p in sorted(source_map): sm.append(f"| {s} | {r} | `{p}` |")
write("_meta/source-map.md", "\n".join(sm + [""]))
print(f"wrote {len(source_map)} generated notes")
